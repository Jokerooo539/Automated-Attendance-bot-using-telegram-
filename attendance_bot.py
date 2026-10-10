#!/usr/bin/env python3
"""
Telegram attendance bot v2.5 -> writes into daily_attendance.xlsx

* Listens in the Block A / B / C Telegram groups.
* Parses messages (both styles: name and code on separate lines, or on the same line).
* Writes the code into the worker's cell for that date. An amount goes on the 2nd line of the same cell.
* Replies in the group with a short summary and anything that needs checking.
* Posts the combined message of all blocks every night to the main group (main_chat) and to the admins.
* Survives network loss: retries by itself, then catches up on every message sent meanwhile (Telegram keeps them 24 h).
* The Excel file lives in the phone's normal storage (config "xlsx_path"), so any Excel app can open and edit it.
* If the file is missing or half-saved by another app, messages are kept and saved automatically once it is fixed.
* Only one copy can run at a time (lock file).
* Admin commands: /file  /summary [dd/mm/yyyy]  /post [dd/mm/yyyy]  /load <A|B|C> + pasted message  /advance Name Amount [date] [code]  /status  /add Name, B, [dd/mm/yyyy]   (anyone: /id)
* From Termux:  python attendance_bot.py add "Name" B [dd/mm/yyyy]

Needs: pip install requests openpyxl
"""
import calendar
import datetime as dt
import json
import logging
import os
import re
import shutil
import sys
import time
from collections import Counter
from copy import copy

import openpyxl
import requests
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment

VERSION = "2.5"
log = logging.getLogger("attendance")
MONTHS = {m: i + 1 for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"])}
OK_WORDS = {"ok", "okay", "yes", "y", "haan", "ha", "han", "confirm", "sahi", "theek"}
MAX_COL = 32            # column AF = day 31 (column B = day 1)
pending = {}            # (chat_id, bot_message_id) -> {"block", "date", "pairs"}
MENU_LABELS = {"📄 Excel file": "/file", "📋 Today's summary": "/summary", "📊 Status": "/status"}
MENU = {"keyboard": [["📄 Excel file", "📋 Today's summary"], ["📊 Status"]], "resize_keyboard": True,
        "is_persistent": True}
BOT_COMMANDS = [("file", "Get the Excel file"), ("summary", "Combined attendance of all blocks (today)"),
                ("advance", "Add an amount: /advance Name 300"), ("add", "Add a worker: /add Name, B"),
                ("load", "Load a block message by hand"), ("status", "Bot status")]


# --------------------------------------------------------------------------- parsing
def parse_date(line):
    m = re.search(r"(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4})", line)
    if m:
        try:
            return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError:
            return None
    m = re.search(r"(\d{1,2})[\s\-]([A-Za-z]{3})[A-Za-z]*[\s\-,]*(\d{4})", line)
    if m and m.group(2).lower() in MONTHS:
        try:
            return dt.date(int(m.group(3)), MONTHS[m.group(2).lower()], int(m.group(1)))
        except ValueError:
            return None
    return None


def _finish(code, rest):
    rest = re.sub(r"(?<=\d),(?=\d{3})", "", rest)
    rest = re.sub(r"RS\.?|₹|/-", " ", rest.upper())
    rest = re.sub(r"[^\w\s+]|_", " ", rest)                # drop emojis and punctuation
    if not re.fullmatch(r"[\d\s+]*", rest):
        return None                      # text after the code -> not a code line
    nums = [int(n) for n in re.findall(r"\d+", rest)]
    return {"code": code, "amount": sum(nums), "amount_text": "+".join(map(str, nums))}


def parse_code_line(raw):
    s = re.sub(r"[*_]", "", raw).strip().upper()
    if not s:
        return None
    s = re.sub(r"^P\s+1(?!\d)", "P1", s)          # "P 1" = P1,  "P 1000" = P + 1000
    m = re.match(r"^(LEFT|PP|P1|P|A)(?![A-Z])\s*(.*)$", s)
    return _finish(m.group(1), m.group(2)) if m else None


def parse_inline(raw):
    s = re.sub(r"[*_]", "", raw).strip()
    m = re.match(r"^(.+?)\s+(LEFT|PP|P\s*1(?!\d)|P|A)(?![A-Za-z])\s*(.*)$", s, re.I)
    if not m:
        return None
    c = parse_code_line(re.sub(r"\s+", "", m.group(2)) + " " + m.group(3))
    return dict(c, name=m.group(1).strip()) if c else None


def make_entry(name, c):
    cell = c["code"] + ("\n" + c["amount_text"] if c["amount"] else "")
    name = re.sub(r"^\s*\d+\s*[.)\-:]\s*", "", re.sub(r"[*_]", "", name)).strip()      # "1. Rahul" -> "Rahul"
    return {"name": name, "code": c["code"], "cell": cell}


def parse_message(text):
    lines = [l.strip() for l in text.splitlines()]
    lines = [l for l in lines if l and l != "."]
    date, entries, skipped, incomplete, dates = None, [], [], [], []
    i = 0
    while i < len(lines):
        line = lines[i]
        i += 1
        d = parse_date(line)
        if d:
            date = d
            dates.append(d)
            continue
        if re.search(r"block\s*[a-c]", line, re.I) or parse_code_line(line):
            continue
        nxt = parse_code_line(lines[i]) if i < len(lines) else None
        if nxt:                                   # two-line style: name, then code
            entries.append(make_entry(line, nxt))
            i += 1
            continue
        inl = parse_inline(line)                  # one-line style: name and code together
        if inl:
            entries.append(make_entry(inl["name"], inl))
            continue
        (incomplete if date else skipped).append(line)
    return {"date": date, "dates": dates, "entries": entries, "skipped": skipped, "incomplete": incomplete}


# --------------------------------------------------------------------------- name matching
def norm(s):
    return re.sub(r"\s+", " ", re.sub(r"[^\w\s]", " ", s.lower())).strip()


def lev(a, b):
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        for j, cb in enumerate(b, 1):
            cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb)))
        prev = cur
    return prev[-1]


def match_names(entries, rows):
    """rows: [{'name','block','row'}] for the block. -> matched, unmatched, ambiguous"""
    table = [dict(r, key=norm(r["name"]), flat=norm(r["name"]).replace(" ", "")) for r in rows]
    matched, unmatched, ambiguous = [], [], []
    for e in entries:
        k = norm(e["name"])
        kf = k.replace(" ", "")

        def dist(r):
            return min(lev(k, r["key"]), lev(kf, r["flat"]),
                       lev(kf, r["flat"][:len(kf)]), lev(kf, r["key"].split(" ")[0]))

        scored = sorted(((dist(r), r["row"], r) for r in table), key=lambda x: (x[0], x[1]))
        if not scored or scored[0][0] > 2:
            first = k.split(" ")[0]
            same = [r for r in table if r["key"].split(" ")[0] == first]
            if same:
                ambiguous.append(dict(e, options=[r["name"] for r in same], reason="surname differs"))
            else:
                unmatched.append(e)
            continue
        ties = [x for x in scored if x[0] == scored[0][0]]
        if len(ties) > 1:
            ambiguous.append(dict(e, options=[t[2]["name"] for t in ties], reason="matches more than one"))
        else:
            matched.append(dict(e, sheet_name=scored[0][2]["name"]))
    return matched, unmatched, ambiguous


# --------------------------------------------------------------------------- Where the files live
class ExcelUnavailable(Exception):
    """The Excel file is missing, can't be opened (e.g. half-saved by another app) or can't be written."""


def data_dir(cfg):
    """The bot's own working files (lock, state, offset). Kept in the bot's folder, not in shared storage."""
    return cfg.get("data_dir") or os.path.dirname(os.path.abspath(__file__))


def backup_dir(cfg):
    return cfg.get("backup_dir") or os.path.join(os.path.dirname(os.path.abspath(cfg["xlsx_path"])), "backups")


def load_book(cfg):
    path = cfg["xlsx_path"]
    if not os.path.exists(path):
        raise ExcelUnavailable("The Excel file is not at %s. Check that Termux has storage access "
                               "(run termux-setup-storage) and that the file was not moved or renamed." % path)
    try:
        return openpyxl.load_workbook(path)
    except Exception as e:
        raise ExcelUnavailable("The Excel file can't be opened (%s). Another app may have left it half-saved. "
                               "Open it again and save, or copy the newest file from the backups folder over it."
                               % type(e).__name__)


def check_storage(cfg):
    """Startup check. Returns a description of the problem, or None if the file is readable and its folder writable."""
    folder = os.path.dirname(os.path.abspath(cfg["xlsx_path"]))
    try:
        os.makedirs(folder, exist_ok=True)
        probe = os.path.join(folder, ".bot_write_test")
        with open(probe, "w") as f:
            f.write("ok")
        os.remove(probe)
    except OSError as e:
        return "I can't write to %s (%s). Run termux-setup-storage in Termux and allow storage access." % (folder, e)
    try:
        load_book(cfg).close()
    except ExcelUnavailable as e:
        return str(e)
    return None


def excel_status(cfg):
    try:
        st = os.stat(cfg["xlsx_path"])
    except OSError:
        return "\nExcel: NOT FOUND at %s" % cfg["xlsx_path"]
    return "\nExcel: %s\nLast saved: %s (%d KB)" % (
        cfg["xlsx_path"], dt.datetime.fromtimestamp(st.st_mtime).strftime("%d-%b %H:%M"), max(1, st.st_size // 1024))


# --------------------------------------------------------------------------- Excel helpers
def read_layout(ws):
    rows, block = [], None
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v is None or not str(v).strip():
            continue
        t = str(v).strip()
        m = re.search(r"block\s*([A-C])", t, re.I)
        if m:
            block = m.group(1).upper()
            continue
        if t.lower() == "name":
            continue
        if block:
            rows.append({"row": r, "name": t, "block": block})
    return rows


def cell_text(ws, r, c):
    v = ws.cell(r, c).value
    return "" if v is None else str(v).strip()


def row_left(ws, r):
    return any(cell_text(ws, r, c).upper().startswith("LEFT") for c in range(2, MAX_COL + 1))


def long_absent(ws, r, col):
    prev = [cell_text(ws, r, c).upper() for c in range(max(2, col - 7), col)]
    return len(prev) >= 7 and all(v.startswith("A") for v in prev)


def ensure_tab(wb, d):
    """Return the sheet for d's month; create it from the latest tab if missing."""
    name = d.strftime("%b").upper()
    if name in wb.sheetnames:
        return wb[name]
    src = wb[wb.sheetnames[-1]]
    ws = wb.copy_worksheet(src)
    ws.title = name
    for rng in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(rng))
    for info in reversed(read_layout(ws)):               # drop workers who have LEFT
        if row_left(ws, info["row"]):
            ws.delete_rows(info["row"])
    ndays = calendar.monthrange(d.year, d.month)[1]
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        is_header = a is not None and str(a).strip().lower() == "name"
        for c in range(2, MAX_COL + 1):
            cell = ws.cell(r, c)
            if is_header:
                day = c - 1
                cell.value = dt.datetime(d.year, d.month, day) if day <= ndays else None
                cell.number_format = "d-mmm"
            else:
                cell.value = None
    log.info("created tab %s", name)
    return ws


def set_cell(ws, r, c, text):
    cell = ws.cell(r, c)
    if isinstance(cell, MergedCell) or any(cell.coordinate in rng for rng in ws.merged_cells.ranges):
        return "merged", ""
    old = cell_text(ws, r, c)
    cell.value = text
    a = cell.alignment
    cell.alignment = Alignment(horizontal=a.horizontal, vertical=a.vertical, wrap_text=True)
    if not old:
        return "ok", ""
    return ("same" if " ".join(old.split()) == " ".join(text.split()) else "changed"), old


def backup_once_a_day(cfg):
    try:
        folder = backup_dir(cfg)
        os.makedirs(folder, exist_ok=True)
        dest = os.path.join(folder, "daily_attendance_%s.xlsx" % dt.date.today().isoformat())
        if not os.path.exists(dest) and os.path.exists(cfg["xlsx_path"]):
            shutil.copy2(cfg["xlsx_path"], dest)
    except OSError as e:
        log.warning("could not make the daily backup: %s", e)


def save_wb(wb, cfg):
    backup_once_a_day(cfg)
    tmp = cfg["xlsx_path"][:-5] + ".tmp.xlsx"
    try:
        wb.save(tmp)                               # write next to the real file, then swap: never a half-written file
        os.replace(tmp, cfg["xlsx_path"])
    except OSError as e:
        try:
            os.remove(tmp)
        except OSError:
            pass
        raise ExcelUnavailable("I couldn't save the Excel file (%s)." % e)


def write_pairs(wb, d, block, pairs):
    """pairs: [(sheet_name, cell_text)] -> list of result dicts"""
    ws = ensure_tab(wb, d)
    col = d.day + 1
    rows = {r["name"]: r for r in read_layout(ws) if r["block"] == block}
    out = []
    for name, text in pairs:
        info = rows.get(name)
        if not info:
            out.append({"name": name, "status": "missing", "old": "", "new": text})
            continue
        status, old = ("left", "") if row_left(ws, info["row"]) and not cell_text(ws, info["row"], col) \
            else set_cell(ws, info["row"], col, text)
        out.append({"name": name, "status": status, "old": old, "new": text})
    return out


def not_reported(ws, d, block, reported, ignore):
    col = d.day + 1
    ign = {norm(n) for n in ignore}
    miss = []
    for r in read_layout(ws):
        if r["block"] != block or r["name"] in reported or norm(r["name"]) in ign:
            continue
        if cell_text(ws, r["row"], col) or row_left(ws, r["row"]) or long_absent(ws, r["row"], col):
            continue
        miss.append(r["name"])
    return miss


# --------------------------------------------------------------------------- message handling
def flat(t):
    return " ".join(t.split())


def process(cfg, block, text, msg_date, edited=False):
    """-> (reply_text or None, pending_dict or None)"""
    p = parse_message(text)
    if not p["entries"] and not p["incomplete"]:
        return None, None
    if len(set(p["dates"])) > 1:
        return ("⚠️ This message has %d different dates (%s). Nothing was saved. "
                "Please post each day as a separate message." % (
                    len(set(p["dates"])), ", ".join(sorted(x.strftime("%d/%m") for x in set(p["dates"]))))), None
    d = p["date"] or msg_date
    note = "" if p["date"] else "\n⚠️ No date in the message, used today's date."
    if abs((d - dt.date.today()).days) > 35:
        return "⚠️ Date %s looks wrong. Nothing was written. Please repost with the correct date." % d.strftime("%d/%m/%Y"), None

    wb = load_book(cfg)
    ws = ensure_tab(wb, d)
    rows = [r for r in read_layout(ws) if r["block"] == block]
    matched, unmatched, ambiguous = match_names(p["entries"], rows)
    results = write_pairs(wb, d, block, [(m["sheet_name"], m["cell"]) for m in matched])
    ws = wb[d.strftime("%b").upper()]
    reported = {m["sheet_name"] for m in matched} | {n for a in ambiguous for n in a["options"]}
    missing = not_reported(ws, d, block, reported, cfg.get("ignore_names", []))
    if any(r["status"] in ("ok", "changed") for r in results):
        save_wb(wb, cfg)

    done = [r for r in results if r["status"] in ("ok", "changed", "same")]
    counts = Counter(m["code"] for m, r in zip(matched, results) if r["status"] in ("ok", "changed", "same"))
    lines = ["✅ Block %s · %s" % (block, d.strftime("%d-%b-%Y")),
             "Saved %d (%s)" % (len(done), ", ".join("%s %d" % kv for kv in sorted(counts.items())) or "none")]
    for r in results:
        if r["status"] == "changed":
            lines.append("🔁 %s: %s → %s" % (r["name"], flat(r["old"]), flat(r["new"])))
        elif r["status"] == "merged":
            lines.append("⛔ %s: that cell is merged in Excel, not saved" % r["name"])
        elif r["status"] == "left":
            lines.append("⛔ %s is marked LEFT, not saved" % r["name"])
        elif r["status"] == "missing":
            lines.append("⛔ %s not found in sheet" % r["name"])
    pend = None
    for a in ambiguous:
        if len(a["options"]) == 1:
            lines.append("❓ %s → %s? (surname differs)" % (a["name"], a["options"][0]))
            pend = pend or {"block": block, "date": d, "pairs": []}
            pend["pairs"].append((a["options"][0], a["cell"]))
        else:
            lines.append("❓ %s could be: %s. Repost with the full name." % (a["name"], ", ".join(a["options"])))
    for u in unmatched:
        lines.append("❓ Name not found: %s" % u["name"])
    for n in p["incomplete"]:
        lines.append("❓ No code after: %s" % n)
    if pend:
        lines.append("Reply \"ok\" to this message to save the ❓ names above.")
    if missing:
        lines.append("Not reported: " + ", ".join(missing))
    if edited and not any(r["status"] == "changed" for r in results) and len(lines) <= 2:
        return None, None
    return "\n".join(lines) + note, pend


def confirm(cfg, pend):
    wb = load_book(cfg)
    res = write_pairs(wb, pend["date"], pend["block"], pend["pairs"])
    if any(r["status"] in ("ok", "changed") for r in res):
        save_wb(wb, cfg)
    return "✅ Saved: " + ", ".join("%s %s" % (r["name"], flat(r["new"])) for r in res
                                    if r["status"] in ("ok", "changed", "same")) \
        + "".join("\n⛔ %s: %s" % (r["name"], r["status"]) for r in res
                  if r["status"] not in ("ok", "changed", "same"))


def combined_with_count(cfg, d):
    """-> (text, number_of_workers_with_an_entry, {'A': n, 'B': n, 'C': n})"""
    counts = {"A": 0, "B": 0, "C": 0}
    wb = load_book(cfg)
    tab = d.strftime("%b").upper()
    if tab not in wb.sheetnames:
        return "No %s tab yet." % tab, 0, counts
    ws, col = wb[tab], d.day + 1
    layout = read_layout(ws)
    lines = [d.strftime("%d-%b-%Y"), ""]
    for b in "ABC":
        lines.append("🔴 Block %s 🔴" % b)
        for r in layout:
            if r["block"] != b:
                continue
            t = flat(cell_text(ws, r["row"], col))
            if t and not t.upper().startswith("LEFT"):
                lines += [r["name"], t, ""]
                counts[b] += 1
    return "\n".join(lines).strip(), sum(counts.values()), counts


def build_combined(cfg, d):
    return combined_with_count(cfg, d)[0]


def load_state(cfg):
    try:
        with open(state_path(cfg)) as f:
            return json.load(f)
    except Exception:
        return {}


def state_path(cfg):
    return os.path.join(data_dir(cfg), "state.json")


def save_state(cfg, st):
    with open(state_path(cfg), "w") as f:
        json.dump(st, f)


def send_long(tg, chat_id, text):
    """Telegram allows ~4096 characters per message, so split long texts on line breaks."""
    chunk = ""
    for line in text.split("\n"):
        if len(chunk) + len(line) + 1 > 3900:
            tg.send(chat_id, chunk.rstrip())
            chunk = ""
        chunk += line + "\n"
    if chunk.strip():
        tg.send(chat_id, chunk.rstrip())


def post_combined(tg, cfg, d, late=False):
    """Send the combined message for date d to the main group and to the admins."""
    main, admins = cfg.get("main_chat"), cfg.get("admins", [])
    if not main and not admins:
        return "set main_chat or admins in config.json"
    text, n, counts = combined_with_count(cfg, d)
    if n == 0:
        if not late:
            for a in admins:
                tg.send(a, "No attendance saved for %s, so nothing was posted." % d.strftime("%d-%b-%Y"))
        return "nothing to post"
    if main:
        send_long(tg, main, text)
    if cfg.get("send_combined_to_admins", True):
        warn = ["⚠️ Block %s: nothing saved for this date" % b for b in "ABC" if counts[b] == 0]
        for a in admins:
            if a != main:
                send_long(tg, a, text + ("\n\n" + "\n".join(warn) if warn else ""))
    return "posted %d workers" % n


def scheduled(tg, cfg, now, st):
    """Daily jobs. State lives in state.json so a restart or an offline night never causes a double post."""
    hhmm_now, changed = now.strftime("%H:%M"), False
    t = cfg.get("combined_post_time", "23:30")
    if t and (cfg.get("main_chat") or cfg.get("admins")):
        done = st.setdefault("posted_days", [])
        if st.get("posted") and st["posted"] not in done:
            done.append(st["posted"])
        for offset in (1, 0):                              # yesterday (missed while offline) and today
            d = now.date() - dt.timedelta(days=offset)
            iso = d.isoformat()
            if iso in done or (offset == 0 and hhmm_now < t) or (offset == 1 and not done):
                continue
            log.info("combined post %s: %s", iso, post_combined(tg, cfg, d, late=(offset == 1)))
            done.append(iso)
            del done[:-10]
            changed = True
    t = cfg.get("daily_file_time")
    today = now.date().isoformat()
    if t and hhmm_now >= t and st.get("file_sent") != today and os.path.exists(cfg["xlsx_path"]):
        for a in cfg.get("admins", []):
            tg.send_file(a, cfg["xlsx_path"])
        st["file_sent"], changed = today, True
    if changed:
        save_state(cfg, st)


def parse_add(rest):
    """'Rohit Kumar, B, 05/10/2026' | 'Rohit Kumar B 05/10/2026' -> (name, block, date) or None"""
    rest = rest.strip()
    if re.search(r"[|,]", rest):
        parts = [p.strip() for p in re.split(r"[|,]", rest) if p.strip()]
        if len(parts) < 2:
            return None
        name, block, date = parts[0], parts[1], (parts[2] if len(parts) > 2 else "")
    else:
        m = re.match(r"^(.+?)\s+(?:block\s*)?([A-Ca-c])(?:\s+(\S+))?$", rest)
        if not m:
            return None
        name, block, date = m.group(1), m.group(2), m.group(3) or ""
    if not re.fullmatch(r"(?i)(?:block\s*)?[A-C]", block):
        return None
    start = parse_date(date) if date else dt.date.today()
    if date and not start:
        return None
    return name.strip(), block[-1].upper(), start


def add_worker(cfg, name, block, start):
    """Insert a new worker row at the end of the block in the start month's tab (and any later tabs)."""
    name = " ".join(name.split())
    if not name or block not in ("A", "B", "C"):
        return "Usage: /add Name, Block (A/B/C), date (optional)"
    wb = load_book(cfg)
    ws = ensure_tab(wb, start)
    layout = read_layout(ws)
    for r in layout:
        if norm(r["name"]) == norm(name):
            return "%s is already in Block %s." % (r["name"].strip(), r["block"])
    targets = wb.worksheets[wb.sheetnames.index(ws.title):]
    for sh in targets:
        if sh.merged_cells.ranges:
            return "The %s tab has merged cells, so I can't insert a row safely. Add this worker by hand." % sh.title
        if not any(r["block"] == block for r in read_layout(sh)):
            return "Block %s not found in the %s tab." % (block, sh.title)
    for sh in targets:
        last = max(r["row"] for r in read_layout(sh) if r["block"] == block)
        sh.insert_rows(last + 1)
        for c in range(1, MAX_COL + 1):
            sh.cell(last + 1, c)._style = copy(sh.cell(last, c)._style)
        sh.row_dimensions[last + 1].height = sh.row_dimensions[last].height
        sh.cell(last + 1, 1).value = name
    save_wb(wb, cfg)
    first = norm(name).split(" ")[0]
    twins = [r["name"].strip() for r in layout if r["block"] == block and norm(r["name"]).split(" ")[0] == first]
    msg = "✅ Added %s to Block %s from %s (%s)." % (
        name, block, start.strftime("%d-%b-%Y"), ", ".join(t.title for t in targets))
    if twins:
        msg += "\n⚠️ Block %s already has %s. Leaders must write the full name for these." % (block, " and ".join(twins))
    return msg


def add_advance(cfg, name, amount, d, code=None):
    """Add an amount under a worker's code for date d, keeping the code and any amount already there."""
    wb = load_book(cfg)
    tab = d.strftime("%b").upper()
    if tab not in wb.sheetnames:
        return "There is no %s tab yet, so I can't add an amount for %s." % (tab, d.strftime("%d-%b-%Y"))
    ws = wb[tab]
    rows = read_layout(ws)
    matched, unmatched, ambiguous = match_names([{"name": name}], rows)
    if ambiguous:
        return "❓ %s could be: %s. Use the full name." % (name, ", ".join(ambiguous[0]["options"]))
    if not matched:
        return "❓ I couldn't find %s in the sheet." % name
    info = next(r for r in rows if r["name"] == matched[0]["sheet_name"])
    col = d.day + 1
    old = [l.strip() for l in cell_text(ws, info["row"], col).splitlines() if l.strip()]
    if old:
        cur = parse_code_line(old[0])
        if not cur:
            return "I can't read the current entry %r for %s. Fix it by hand first." % (" ".join(old), info["name"].strip())
        code = cur["code"]
        parts = re.findall(r"\d+", " ".join(old[1:]))
    else:
        if not code:
            return ("%s has nothing saved for %s, so I don't know the code. Add it, for example: "
                    "/advance %s %s P1" % (info["name"].strip(), d.strftime("%d-%b-%Y"), name, amount))
        parts = []
    parts += re.findall(r"\d+", amount)
    cell = code + "\n" + "+".join(parts)
    status, _ = set_cell(ws, info["row"], col, cell)
    if status == "merged":
        return "⛔ That cell is merged in Excel, so I can't change it."
    save_wb(wb, cfg)
    return "✅ %s · %s: %s" % (info["name"].strip(), d.strftime("%d-%b-%Y"), cell.replace("\n", " "))


def advance_command(cfg, args):
    """/advance Bablu 300 [dd/mm/yyyy] [P1]"""
    d, code, amount, name = None, None, None, []
    for a in args:
        if re.search(r"[/.\-]", a) and parse_date(a):
            d = parse_date(a)
        elif re.fullmatch(r"(?i)p1|pp|p|a", a):
            code = a.upper()
        elif re.fullmatch(r"\d+(\+\d+)*", a) and amount is None:
            amount = a
        else:
            name.append(a)
    if not name or not amount:
        return "Usage: /advance Name Amount [date] [code]\nExample: /advance Bablu 300 02/10/2026"
    return add_advance(cfg, " ".join(name), amount, d or dt.date.today(), code)


# --------------------------------------------------------------------------- Telegram
class Tg:
    def __init__(self, token):
        self.api = "https://api.telegram.org/bot%s/" % token

    def call(self, method, **params):
        r = requests.post(self.api + method, json=params, timeout=(10, 45))
        return r.json()

    def send(self, chat_id, text, reply_to=None, reply_markup=None):
        extra = {"reply_to_message_id": reply_to, "allow_sending_without_reply": True} if reply_to else {}
        if reply_markup:
            extra["reply_markup"] = reply_markup
        r = self.call("sendMessage", chat_id=chat_id, text=text, **extra)
        return (r.get("result") or {}).get("message_id")

    def send_file(self, chat_id, path):
        with open(path, "rb") as f:
            requests.post(self.api + "sendDocument", data={"chat_id": chat_id},
                          files={"document": f}, timeout=(10, 120))


def handle_command(tg, cfg, msg, text):
    chat_id = msg["chat"]["id"]
    uid = (msg.get("from") or {}).get("id")
    cmd = text.split()[0].split("@")[0].lower()
    args = text.split()[1:]
    if cmd == "/id":
        tg.send(chat_id, "chat id: %s\nyour user id: %s" % (chat_id, uid))
        return
    if uid not in cfg.get("admins", []):
        return
    if msg["chat"].get("type", "private") != "private":
        tg.send(chat_id, "Please send admin commands in a private chat with me, not in a group.")
        return
    if cmd in ("/start", "/menu"):
        tg.send(chat_id, "Use the buttons below to get the Excel file or today's attendance.", reply_markup=MENU)
    elif cmd == "/file":
        if not os.path.exists(cfg["xlsx_path"]):
            raise ExcelUnavailable("The Excel file is not at %s." % cfg["xlsx_path"])
        tg.send_file(chat_id, cfg["xlsx_path"])
    elif cmd == "/summary":
        d = parse_date(args[0]) if args else dt.date.today()
        tg.send(chat_id, build_combined(cfg, d or dt.date.today()))
    elif cmd == "/advance":
        tg.send(chat_id, advance_command(cfg, args))
    elif cmd == "/post":
        d = (parse_date(args[0]) if args else None) or dt.date.today()
        tg.send(chat_id, post_combined(tg, cfg, d))
    elif cmd == "/load":                    # /load B  (new line)  <pasted attendance message>
        head, _, body = text.partition("\n")
        blk = (head.split()[1].upper() if len(head.split()) > 1 else "")
        if blk not in ("A", "B", "C") or not body.strip():
            tg.send(chat_id, "Usage: /load B, then a new line, then paste the block's attendance message.")
        else:
            reply, pend = process(cfg, blk, body, dt.date.today())
            mid = tg.send(chat_id, reply or "I couldn't find any attendance lines in that message.")
            if pend and mid:
                pending[(chat_id, mid)] = pend
    elif cmd == "/add":
        parsed = parse_add(text.split(None, 1)[1]) if len(text.split(None, 1)) > 1 else None
        tg.send(chat_id, add_worker(cfg, *parsed) if parsed else
                "Usage: /add Name, Block (A/B/C), date (optional)\nExample: /add Rohit Kumar, B, 05/10/2026")
    elif cmd == "/status":
        tg.send(chat_id, "Bot v%s running. Groups: %s\nNightly post: %s" % (
            VERSION, ", ".join("%s=%s" % (v, k) for k, v in cfg["block_chats"].items()),
            ("%s to %s" % (cfg.get("combined_post_time"), cfg["main_chat"]))
            if cfg.get("main_chat") and cfg.get("combined_post_time") else "off (set main_chat and combined_post_time)")
            + excel_status(cfg))


def handle_update(tg, cfg, upd):
    msg = upd.get("message") or upd.get("edited_message")
    if not msg:
        return
    edited = "edited_message" in upd
    chat_id = msg["chat"]["id"]
    text = msg.get("text") or msg.get("caption") or ""
    text = MENU_LABELS.get(text.strip(), text)
    if text.startswith("/"):
        try:
            handle_command(tg, cfg, msg, text)
        except ExcelUnavailable as e:
            tg.send(chat_id, "⚠️ " + str(e))
        return
    # confirmation of a pending "❓" list
    rep = (msg.get("reply_to_message") or {}).get("message_id")
    if rep and (chat_id, rep) in pending and text.strip().lower() in OK_WORDS:
        reply = confirm(cfg, pending[(chat_id, rep)])      # if the file is unavailable this raises and stays pending
        pending.pop((chat_id, rep), None)
        tg.send(chat_id, reply, reply_to=msg["message_id"])
        return
    block = cfg["block_chats"].get(str(chat_id))
    if not block:
        return
    msg_date = dt.datetime.fromtimestamp(msg["date"]).date()
    reply, pend = process(cfg, block, text, msg_date, edited)
    if reply:
        mid = tg.send(chat_id, reply, reply_to=msg["message_id"])
        if pend and mid:
            pending[(chat_id, mid)] = pend


def acquire_lock(cfg):
    """Only one copy of the bot may run (two copies polling Telegram fight each other)."""
    try:
        import fcntl
    except ImportError:                       # e.g. Windows
        return None
    f = open(os.path.join(data_dir(cfg), "bot.lock"), "w")
    try:
        fcntl.flock(f, fcntl.LOCK_EX | fcntl.LOCK_NB)
    except OSError:
        sys.exit("Another copy of the bot is already running. Stop it first: pkill -f attendance_bot.py")
    f.write(str(os.getpid()))
    f.flush()
    return f


def tell_admins(tg, cfg, text):
    for a in cfg.get("admins", []):
        try:
            tg.send(a, text)
        except Exception:
            log.exception("could not notify admin")


def step(tg, cfg, ctx, now=None):
    """One polling round. Returns seconds to sleep afterwards. Network errors propagate to run()."""
    now = now or dt.datetime.now()
    params = {"timeout": 30, "allowed_updates": ["message", "edited_message"]}
    if ctx.get("offset"):
        params["offset"] = ctx["offset"]
    resp = tg.call("getUpdates", **params)
    if not resp.get("ok"):                           # wrong/revoked token, webhook conflict, etc.
        log.error("Telegram refused getUpdates: %s", resp.get("description"))
        return 15
    updates = resp.get("result", [])
    came_back = ctx.pop("offline_since", None)
    ctx["fails"] = 0
    for upd in updates:                              # messages sent while offline arrive here, oldest first
        try:
            handle_update(tg, cfg, upd)
        except ExcelUnavailable as e:                # keep this message (offset not advanced); Telegram resends it
            log.error("Excel unavailable, will retry: %s", e)
            if time.time() - ctx.get("excel_alert", 0) > 1800:
                ctx["excel_alert"] = time.time()
                tell_admins(tg, cfg, "⚠️ I can't save attendance right now. %s\nNothing is lost: the messages are kept "
                                     "and saved automatically once this is fixed (Telegram holds them for 24 hours)." % e)
            return 30
        except Exception:
            log.exception("failed on update %s", upd.get("update_id"))
        ctx["offset"] = upd["update_id"] + 1
        with open(ctx["offset_file"], "w") as f:
            f.write(str(ctx["offset"]))
    if updates and ctx.get("excel_alert"):
        ctx["excel_alert"] = 0
        tell_admins(tg, cfg, "✅ The Excel file works again. The waiting messages have been saved.")
    notice = None
    if came_back and (updates or time.time() - came_back >= 120):
        notice = "🔌 Back online after %d min. Caught up on %d message(s) that arrived meanwhile." % (
            max(1, int((time.time() - came_back) / 60)), len(updates))
    elif ctx.get("first") and updates:
        notice = "🔄 Bot started. Caught up on %d message(s) that arrived while it was off." % len(updates)
    ctx["first"] = False
    if notice:
        log.info(notice)
        tell_admins(tg, cfg, notice)
    if len(updates) >= 100:                          # more are waiting; finish them before the daily jobs
        return 0
    try:
        scheduled(tg, cfg, now, ctx["state"])
    except ExcelUnavailable as e:
        log.warning("daily jobs postponed: %s", e)
    if time.time() - ctx.get("hb", 0) > 60:          # "last seen alive" marker for long-outage warnings
        ctx["state"]["alive"] = now.isoformat(timespec="seconds")
        save_state(cfg, ctx["state"])
        ctx["hb"] = time.time()
    return 0


class TokenFilter(logging.Filter):
    """Masks the bot token in every log line, including tracebacks (error texts can contain the request URL)."""
    def __init__(self, token):
        super().__init__()
        self.token = token

    def filter(self, record):
        record.msg = str(record.getMessage()).replace(self.token, "<token>")
        record.args = ()
        if record.exc_info:
            record.exc_text = logging.Formatter().formatException(record.exc_info).replace(self.token, "<token>")
        return True


def run(cfg):
    for h in logging.getLogger().handlers:
        h.addFilter(TokenFilter(cfg["token"]))
    lock = acquire_lock(cfg)                         # keep the reference so the lock stays held
    tg = Tg(cfg["token"])
    offset_file = os.path.join(data_dir(cfg), "offset.txt")
    ctx = {"offset": int(open(offset_file).read()) if os.path.exists(offset_file) else None,
           "state": load_state(cfg), "first": True, "offset_file": offset_file}
    log.info("bot v%s started", VERSION)
    try:
        tg.call("setMyCommands", scope={"type": "all_private_chats"},
                commands=[{"command": c, "description": d} for c, d in BOT_COMMANDS])
    except Exception as e:
        log.warning("could not register the command list: %s", e)
    problem = check_storage(cfg)
    if problem:
        log.error("storage problem: %s", problem)
        tell_admins(tg, cfg, "⚠️ " + problem)
    alive = ctx["state"].get("alive")
    if alive:
        hours = (dt.datetime.now() - dt.datetime.fromisoformat(alive)).total_seconds() / 3600
        if hours >= 20:
            tell_admins(tg, cfg, "⚠️ The bot was off for about %d hours. Telegram keeps messages for only 24 hours, "
                                 "so older attendance may be missing. Check the sheet and use /load for anything absent." % hours)
    while True:
        try:
            delay = step(tg, cfg, ctx)
            if delay:
                time.sleep(delay)
        except requests.RequestException as e:       # no internet, DNS failure, timeout
            if "offline_since" not in ctx:
                ctx["offline_since"] = time.time()
                log.warning("network problem, will keep retrying: %s", e)
            ctx["fails"] = ctx.get("fails", 0) + 1
            time.sleep(min(10 * ctx["fails"], 60))
        except Exception:
            log.exception("loop error")
            time.sleep(10)


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s",
                        handlers=[logging.StreamHandler(), logging.FileHandler("bot.log")])
    args = sys.argv[1:]
    if args and args[0] == "add":                      # python attendance_bot.py add "Name" B [dd/mm/yyyy]
        parsed = parse_add(" ".join(args[1:]))
        if not parsed:
            sys.exit('Usage: python attendance_bot.py add "Rohit Kumar" B [dd/mm/yyyy]')
        try:
            print(add_worker(json.load(open("config.json")), *parsed))
        except ExcelUnavailable as e:
            sys.exit(str(e))
    else:
        run(json.load(open(args[0] if args else "config.json")))
